"""Assisted workflow on the existing application database and document renderer."""

from __future__ import annotations

import hashlib
import json
import re
import threading
import uuid
from pathlib import Path

from PyPDF2 import PdfReader

from core.assisted_matching import evaluate, select_entries, use_ai_selection
from core.job_discovery import canonical_url, discover, now, verify_url
from core.resume_renderer import render_resume_to_pdf

STATES = {
    "saved",
    "skipped",
    "selected",
    "resume_prepared",
    "applied",
    "response_received",
    "interview",
    "rejected",
    "accepted",
}
_search_lock = threading.Lock()
_resume_lock = threading.Lock()


def encode(data):
    return json.dumps(data, ensure_ascii=False, sort_keys=True)


def persist_job(db, job, analysis):
    """Atomic dedup by provider ID, canonical URL, or company/title/location."""
    with db._connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute(
            """SELECT a.id, j.payload, j.discovered_at FROM applications a
            LEFT JOIN discovered_jobs j ON j.application_id=a.id
            WHERE (a.external_id=? AND a.platform=?) OR j.canonical_url=? OR j.fingerprint=?
            LIMIT 1""",
            (job["external_id"], job["source"], job["canonical_url"], job["fingerprint"]),
        ).fetchone()
        stamp = now()
        if row:
            app_id = row["id"]
            first_seen = row["discovered_at"] or stamp
            if row["payload"]:
                previous = json.loads(row["payload"])
                if (
                    previous["description"] == job["description"]
                    and previous["apply_url"] == job["apply_url"]
                ):
                    for k in ("url_status", "url_message", "checked_at", "availability"):
                        if k in previous:
                            job[k] = previous[k]
        else:
            cur = conn.execute(
                """INSERT INTO applications(external_id, platform, job_title, company, location,
                apply_url, match_score, status) VALUES(?,?,?,?,?,?,?, 'saved')""",
                (
                    job["external_id"],
                    job["source"],
                    job["title"],
                    job["company"],
                    job["location"],
                    job["apply_url"],
                    analysis["score"],
                ),
            )
            app_id, first_seen = cur.lastrowid, stamp
        job["discovered_at"], job["match"] = first_seen, analysis
        conn.execute(
            """INSERT INTO discovered_jobs VALUES(?,?,?,?,?,?) ON CONFLICT(application_id)
            DO UPDATE SET canonical_url=excluded.canonical_url, fingerprint=excluded.fingerprint,
            payload=excluded.payload, last_seen_at=excluded.last_seen_at""",
            (app_id, job["canonical_url"], job["fingerprint"], encode(job), first_seen, stamp),
        )
        conn.execute(
            "UPDATE applications SET match_score=?, apply_url=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (analysis["score"], job["apply_url"], app_id),
        )
        return app_id


def list_jobs(db):
    with db._connect() as conn:
        rows = conn.execute("""SELECT j.payload, a.id, a.status FROM discovered_jobs j
            JOIN applications a ON a.id=j.application_id ORDER BY a.match_score DESC, j.last_seen_at DESC""").fetchall()
        jobs = []
        for row in rows:
            job = json.loads(row["payload"])
            job.update(id=row["id"], status=row["status"])
            job["resumes"] = [
                dict(r)
                for r in conn.execute(
                    """SELECT v.id, v.created_at, v.llm_provider FROM resume_versions v
                JOIN assisted_resumes r ON r.version_id=v.id WHERE v.application_id=? ORDER BY v.id DESC""",
                    (row["id"],),
                )
            ]
            jobs.append(job)
        return jobs


def get_job(db, job_id):
    return next((j for j in list_jobs(db) if j["id"] == job_id), None)


def save_payload(db, job):
    with db._connect() as conn:
        conn.execute(
            "UPDATE discovered_jobs SET payload=?, canonical_url=? WHERE application_id=?",
            (encode(job), canonical_url(job["apply_url"]), job["id"]),
        )
        conn.execute(
            "UPDATE applications SET apply_url=?,match_score=? WHERE id=?",
            (job["apply_url"], job["match"]["score"], job["id"]),
        )


def search_jobs(db, config):
    if not _search_lock.acquire(blocking=False):
        raise ValueError("Já existe uma busca em andamento")
    try:
        jobs, errors = discover(config.search_criteria)
        for message in errors:
            db.save_feed_event("error", message=message)
        criteria = config.search_criteria
        filtered = []
        for job in jobs:
            title, text = job["title"].lower(), job["description"].lower()
            if criteria.job_titles and not any(t.lower() in title for t in criteria.job_titles):
                continue
            if criteria.keywords_include and not any(
                k.lower() in text + " " + title for k in criteria.keywords_include
            ):
                continue
            analysis = evaluate(config.profile, criteria, job, config.company_blacklist)
            # Exclusions are explicit. Unknown metadata is never a rejection.
            if any(
                x in analysis["blockers"]
                for x in (
                    "Palavra-chave excluída encontrada",
                    "Empresa bloqueada",
                    "Modalidade incompatível com remoto",
                )
            ):
                continue
            filtered.append((job, analysis))
        filtered.sort(key=lambda x: x[1]["score"], reverse=True)
        ids = []
        for job, analysis in filtered[: criteria.max_results]:
            ids.append(persist_job(db, job, analysis))
        # Bound verification cost. Other results are explicitly unverified.
        for job_id in list(dict.fromkeys(ids))[:5]:
            job = get_job(db, job_id)
            if job["url_status"] != "verified":
                job.update(verify_url(job))
                save_payload(db, job)
        db.save_feed_event(
            "info",
            message=f"Busca ATS: {len(jobs)} registros recebidos; {len(set(ids))} vagas no resultado; {len(errors)} falhas de fonte",
        )
        return dict(
            job_ids=list(dict.fromkeys(ids)),
            received=len(jobs),
            matching=len(filtered),
            errors=errors,
        )
    finally:
        _search_lock.release()


def _safe(value):
    value = re.sub(r"[^a-zA-Z0-9_-]+", "-", value).strip(".-")[:65]
    return (
        value
        if value and value.upper() not in {"CON", "PRN", "AUX", "NUL", "COM1", "LPT1"}
        else "company"
    )


def _literal(text):
    # Keep user text readable while preventing it from injecting Markdown layout.
    return str(text).replace("\r", " ").replace("\n", " ").replace("**", "").strip()


def build_markdown(profile, entries, language):
    pt = language == "pt-BR"
    headings = {
        "experience": "Experiência profissional" if pt else "Experience",
        "education": "Formação acadêmica" if pt else "Education",
        "project": "Projetos" if pt else "Projects",
        "skill": "Competências" if pt else "Skills",
        "language": "Idiomas" if pt else "Languages",
        "certification": "Certificações" if pt else "Certifications",
        "research": "Pesquisa e atividades acadêmicas"
        if pt
        else "Research and academic activities",
        "other": "Outras atividades" if pt else "Additional information",
    }
    contacts = [
        profile.email,
        profile.phone_full,
        profile.location,
        profile.linkedin_url,
        profile.github_url,
        profile.portfolio_url,
        *profile.professional_links,
    ]
    lines = ["# " + _literal(profile.full_name), " | ".join(_literal(x) for x in contacts if x), ""]
    for category, heading in headings.items():
        group = [
            e
            for e in entries
            if e["category"] == category or (category == "other" and e["category"] not in headings)
        ]
        if not group:
            continue
        lines.append("## " + heading)
        for entry in group:
            header = " - ".join(
                _literal(entry[k]) for k in ("organization", "title") if entry.get(k)
            )
            if header:
                lines.append("### " + header)
                dates = " - ".join(
                    _literal(entry[k]) for k in ("start_date", "end_date") if entry.get(k)
                )
                details = " - ".join(
                    x
                    for x in (
                        dates,
                        _literal(entry.get("location", "")),
                        _literal(entry.get("experience_type", "")),
                    )
                    if x
                )
                if details:
                    lines.append(details)
            lines.append("- " + _literal(entry["text"]))
            if entry.get("technologies"):
                lines.append(
                    ("Tecnologias: " if pt else "Technologies: ") + ", ".join(entry["technologies"])
                )
            if entry.get("proficiency"):
                lines.append(_literal(entry["proficiency"]))
            if entry.get("url"):
                lines.append(_literal(entry["url"]))
            lines.append("")
    return "\n".join(lines)


def validate_pdf(path, profile, expected_claims=()):
    if not path.exists() or path.stat().st_size < 500:
        raise ValueError("PDF vazio ou ausente")
    reader = PdfReader(str(path))
    text = "\n".join(p.extract_text() or "" for p in reader.pages)
    normalized = " ".join(text.split())
    for claim in expected_claims:
        if " ".join(claim.split()) not in normalized:
            raise ValueError(
                "PDF perdeu texto de uma informação confirmada; verifique caracteres e formatação"
            )
    for value in (profile.full_name, profile.email, profile.phone_full):
        if value and " ".join(value.split()) not in normalized:
            raise ValueError("PDF não contém nome ou contato obrigatório")
    if len(normalized) < 70 or re.search(
        r"\[(?:full name|name|company|degree|year)\]|\b(?:TODO|PLACEHOLDER|LOREM IPSUM)\b",
        text,
        re.I,
    ):
        raise ValueError("PDF vazio ou com placeholders")
    return dict(pages=len(reader.pages), extracted_characters=len(text), valid=True)


def generate_resume(db, config, job, data_dir, use_ai=False, force=False):
    if not config.profile.full_name.strip() or not config.profile.email.strip():
        raise ValueError("Preencha nome e e-mail no perfil")
    entries = select_entries(config.profile, job)
    if not entries:
        raise ValueError("Confirme ao menos uma informação profissional no perfil")
    language = config.search_criteria.resume_language
    if language == "auto":
        language = (
            "pt-BR"
            if job.get("language") in ("pt", "pt-BR")
            or re.search(r"\b(você|requisitos|experiência)\b", job["description"], re.I)
            else "en"
        )
    profile_input = config.profile.model_dump()
    profile_input["career_entries"] = [e for e in profile_input["career_entries"] if e["confirmed"]]
    cache_input = dict(
        profile=profile_input,
        description=job["description"],
        language=language,
        ai=use_ai,
        model=config.llm.model if use_ai else "",
        provider=config.llm.provider if use_ai else "",
        endpoint=config.llm.endpoint if use_ai else "",
        renderer=3,
    )
    cache_key = hashlib.sha256(encode(cache_input).encode()).hexdigest()
    with _resume_lock:
        with db._connect() as conn:
            cached = conn.execute(
                """SELECT v.id,v.resume_pdf_path,r.docx_path FROM assisted_resumes r
                JOIN resume_versions v ON v.id=r.version_id WHERE v.application_id=? AND r.cache_key=? ORDER BY v.id DESC LIMIT 1""",
                (job["id"], cache_key),
            ).fetchone()
        if cached and not force and Path(cached["docx_path"]).exists():
            try:
                validation = validate_pdf(Path(cached["resume_pdf_path"]), config.profile)
                return dict(version_id=cached["id"], cached=True, validation=validation)
            except (ValueError, OSError):
                pass
        if use_ai:
            entries = use_ai_selection(entries, job, config.llm)
        md = build_markdown(config.profile, entries, language)
        folder = (
            data_dir
            / "outputs"
            / _safe(job["company"])
            / str(job["id"])
            / (now()[:10] + "-" + uuid.uuid4().hex[:10])
        )
        folder.mkdir(parents=True, exist_ok=True)
        pdf_path, md_path, docx_path = (
            folder / "resume.pdf",
            folder / "resume.md",
            folder / "resume.docx",
        )
        md_path.write_text(md, encoding="utf-8")
        render_resume_to_pdf(md, pdf_path)
        validation = validate_pdf(pdf_path, config.profile, [_literal(e["text"]) for e in entries])
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor

        doc = Document()
        sec = doc.sections[0]
        sec.top_margin = sec.bottom_margin = Inches(0.6)
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10)
        for style_name in ("Title", "Heading 1", "Heading 2"):
            doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
        # Some Word templates carry an inherited blue border in the Title style.
        for root in (doc._element, doc.styles.element):
            for border in root.xpath(".//w:pBdr"):
                border.getparent().remove(border)
        for line in md.splitlines():
            if line.startswith("# "):
                p = doc.add_paragraph(line[2:], style="Title")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=1)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=2)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line:
                doc.add_paragraph(line)
        doc.save(docx_path)
        provenance = dict(
            profile_snapshot=config.profile.model_dump(),
            selected_facts=entries,
            job_id=job["id"],
            description_hash=hashlib.sha256(job["description"].encode()).hexdigest(),
            validation=validation,
            language=language,
            content_language="original confirmed text",
            ai_selection=use_ai,
        )
        provenance_path = folder / "provenance.json"
        provenance_path.write_text(encode(provenance), encoding="utf-8")
        (folder / "job_description.txt").write_text(job["description"], encoding="utf-8")
        (folder / "match_analysis.json").write_text(encode(job["match"]), encoding="utf-8")
        version_id = db.save_resume_version(
            job["id"],
            job["title"],
            job["company"],
            str(md_path),
            str(pdf_path),
            job["match"]["score"],
            config.llm.provider if use_ai else "local",
            config.llm.model if use_ai else "",
            "confirmed_profile",
            encode([e["id"] for e in entries]),
        )
        with db._connect() as conn:
            conn.execute(
                "INSERT INTO assisted_resumes VALUES(?,?,?,?)",
                (version_id, cache_key, str(docx_path), str(provenance_path)),
            )
            conn.execute(
                "UPDATE applications SET resume_path=?, status=CASE WHEN status IN ('saved','selected') THEN 'resume_prepared' ELSE status END WHERE id=?",
                (str(pdf_path), job["id"]),
            )
        db.save_feed_event(
            "info",
            job_title=job["title"],
            company=job["company"],
            message="Currículo validado e registrado",
        )
        return dict(version_id=version_id, cached=False, validation=validation)
