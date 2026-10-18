"""Fictional document with the same structure as the reported import incident."""

import io
import json
from collections import Counter
from pathlib import Path

import pytest

from config.settings import AppConfig, load_config, save_config
from core.ai_engine import AIError
from core.assisted_matching import professional_months, use_ai_selection
from core.assisted_workflow import build_markdown
from core.document_parser import extract_import_text
from core.profile_import import basic_text_draft, draft_profile, intelligent_draft
from core.profile_schema import deduplicate

FIXTURE = Path(__file__).parent / "fixtures" / "structured_career.md"


@pytest.fixture
def document():
    return FIXTURE.read_text(encoding="utf-8")


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("AUTOAPPLY_DATA_DIR", str(tmp_path))
    monkeypatch.setattr("config.settings._check_keyring", lambda: False)
    import app_state
    from app import app
    from db.database import Database

    monkeypatch.setattr(app_state, "db", Database(tmp_path / "test.db"))
    config = AppConfig(
        profile=dict(
            first_name="Existing",
            last_name="Person",
            email="existing@example.invalid",
            phone="",
            city="",
            state="",
            bio="",
        ),
        search_criteria=dict(job_titles=[], locations=[]),
    )
    save_config(config)
    return app.test_client()


def upload(client, text, **options):
    return client.post(
        "/api/assistant/import",
        data={"file": (io.BytesIO(text.encode()), "fictional.md"), **options},
    )


def test_structured_fields_not_raw_markdown(document):
    result = draft_profile(document, "fictional.md")
    assert Counter(e["category"] for e in result["entries"]) == dict(
        experience=2, education=1, project=7, skill=4, language=3
    )
    assert set(result["contacts"]) == {
        "first_name",
        "last_name",
        "email",
        "phone_country_code",
        "phone",
        "city",
        "state",
        "country",
        "github_url",
        "linkedin_url",
        "portfolio_url",
    }
    assert result["contacts"]["first_name"] == "Pessoa"
    assert result["contacts"]["phone_country_code"] == "+55"
    assert result["contacts"]["portfolio_url"] == "https://portfolio.example.invalid"
    assert all(not e["confirmed"] and e["source_excerpt"] for e in result["entries"])
    assert all(
        e["import_mode"] == "structured"
        and e["import_engine"] == "structured-local"
        and e["import_version"] == 3
        for e in result["entries"]
    )
    assert not any(
        "**Nome profissional:**" in e["text"] or "example.invalid" in e["text"]
        for e in result["entries"]
    )
    skills = [
        tech.casefold()
        for entry in result["entries"]
        if entry["category"] == "skill"
        for tech in entry["technologies"]
    ]
    assert len(skills) == len(set(skills))


def test_overlapping_employments_and_nested_activity(document):
    entries = draft_profile(document, "fixture.md")["entries"]
    jobs = [e for e in entries if e["category"] == "experience"]
    assert [(j["organization"], j["start_date"], j["end_date"]) for j in jobs] == [
        ("Empresa Aurora", "2025-02", "presente"),
        ("Estúdio Boreal", "2026-08", "presente"),
    ]
    assert len(jobs[0]["activities"]) == 2 and not jobs[1]["activities"]
    assert "julho" in jobs[0]["activities"][1]["end_date"]
    assert "setembro de 2026" in " ".join(jobs[1]["date_notes"])
    assert professional_months(jobs) == professional_months(jobs[:1])
    assert not any(e["organization"] == "Empresa Frutal" for e in jobs)


def test_uncertainties_language_and_confidentiality_survive_without_resume_leak(document):
    result = draft_profile(document, "fixture.md")
    languages = {e["title"]: e for e in result["entries"] if e["category"] == "language"}
    assert (
        languages["Inglês"]["proficiency"] == "B2"
        and languages["Inglês"]["proficiency_basis"] == "autoavaliado"
    )
    assert (
        languages["Japonês"]["proficiency"] == "N5" and "aproximado" in languages["Japonês"]["text"]
    )
    assert not any(e["category"] == "certification" for e in result["entries"])
    edu = next(e for e in result["entries"] if e["category"] == "education")
    assert "2028" in edu["end_date"] and "sem mês" in edu["end_date"]
    assert "2027" in " ".join(edu["date_notes"])
    assert result["location_is_reference"] is True
    assert any("Projeto Sigiloso" in c["text"] for c in result["constraints"])
    profile = AppConfig(
        profile=dict(**result["contacts"], bio="", location_is_reference=True),
        search_criteria=dict(job_titles=[], locations=[]),
    ).profile
    markdown = build_markdown(profile, result["entries"], "pt-BR")
    assert "Projeto Sigiloso" not in markdown and "Não afirmar" not in markdown
    assert "Cidade Exemplo" not in markdown and "não cofundador" not in markdown
    assert "graduação em andamento" in markdown
    assert "prova de conceito" in markdown.lower()


def test_idempotent_identity_preserves_edits_and_distinct_periods(document):
    first = draft_profile(document, "a.md")["entries"]
    repeated = draft_profile(document, "renamed.md")["entries"]
    assert [e["id"] for e in first] == [e["id"] for e in repeated]
    first[0]["text"] = "User correction is preserved."
    assert deduplicate(repeated, first) == ([], 17)
    changed = {**repeated[0], "start_date": "2020-02"}
    changed.pop("import_key")
    changed.pop("id")
    assert len(deduplicate([changed], first)[0]) == 1


def test_basic_mode_never_inserts_raw_facts(document, client):
    result = basic_text_draft(document, "basic.md")
    assert result["mode"] == "basic" and result["text"] == document
    assert result["entries"] == [] and not result["contacts"]
    response = upload(client, document, mode="basic")
    assert response.status_code == 200 and not response.json["entries"]


def test_import_review_does_not_apply_until_save_and_repeat_is_empty(client, document, tmp_path):
    before = load_config().profile.model_dump()
    response = upload(client, document, mode="structured")
    assert response.status_code == 200 and response.json["mode"] == "structured"
    assert load_config().profile.model_dump() == before
    draft = response.json
    new_profile = {
        **draft["contacts"],
        "career_entries": draft["entries"],
        "generation_constraints": draft["constraints"],
        "contact_notes": draft["contact_notes"],
        "location_is_reference": draft["location_is_reference"],
    }
    assert client.put("/api/assistant/config", json={"profile": new_profile}).status_code == 200
    assert all(not e.confirmed for e in load_config().profile.career_entries)
    assert len(load_config().profile.generation_constraints) > 0
    assert list((tmp_path / "backups").glob("profile-before-import-*.json"))
    again = upload(client, document, mode="structured").json
    assert again["entries"] == [] and again["duplicate_count"] == 17


@pytest.mark.parametrize("extension", [".md", ".txt", ".docx", ".pdf"])
def test_all_document_formats_preserve_sections(extension, document, tmp_path):
    path = tmp_path / ("fixture" + extension)
    if extension in (".md", ".txt"):
        path.write_text(document, encoding="utf-8")
    elif extension == ".docx":
        from docx import Document

        doc = Document()
        for line in document.splitlines():
            if line.startswith("#"):
                depth = len(line) - len(line.lstrip("#"))
                doc.add_heading(line.lstrip("# "), level=max(0, depth - 1))
            else:
                doc.add_paragraph(line)
        doc.save(path)
    else:
        from reportlab.pdfgen import canvas

        pdf = canvas.Canvas(str(path))
        y = 800
        for line in document.splitlines():
            if y < 30:
                pdf.showPage()
                y = 800
            pdf.setFont("Helvetica", 6)
            pdf.drawString(20, y, line)
            y -= 10
        pdf.save()
    extracted = extract_import_text(path)
    result = draft_profile(extracted, path.name)
    assert Counter(e["category"] for e in result["entries"])["project"] == 7
    assert Counter(e["category"] for e in result["entries"])["experience"] == 2
    assert result["contacts"]["email"] == "pessoa@example.invalid"


def test_intelligent_recognized_markdown_does_not_need_paid_call(document, monkeypatch):
    monkeypatch.setattr(
        "core.ai_engine.invoke_llm",
        lambda *a, **k: pytest.fail("Recognized document must not require an AI call"),
    )
    from config.settings import LLMConfig

    result = intelligent_draft(document, "structured.md", LLMConfig())
    assert result["mode"] == "ai" and result["ai_used"] is False and len(result["entries"]) == 17


def test_safe_ai_failure_no_fallback_or_profile_mutation(client, monkeypatch):
    before = load_config().profile.model_dump()

    def fail(*a, **k):
        raise AIError(
            "permission", "Projeto sem acesso ao modelo.", status=403, code="permission_denied"
        )

    monkeypatch.setattr("core.ai_engine.invoke_llm", fail)
    response = upload(client, "Some unstructured career text", mode="ai", consent_ai="true")
    assert response.status_code == 422
    assert response.json["error"]["category"] == "permission"
    assert "entries" not in response.json and load_config().profile.model_dump() == before
    assert upload(client, "text", mode="ai").status_code == 400


def test_ai_schema_and_source_validation(monkeypatch):
    from config.settings import LLMConfig

    captured = {}

    def invoke(prompt, llm, **kwargs):
        captured.update(kwargs)
        return json.dumps(
            {
                "entries": [
                    dict(
                        category="project",
                        title="Example",
                        text="Built an API.",
                        source_excerpt="Example: Built an API.",
                    )
                ]
            }
        )

    monkeypatch.setattr("core.ai_engine.invoke_llm", invoke)
    result = intelligent_draft("Example: Built an API.", "fixture.txt", LLMConfig())
    assert result["ai_used"] and result["entries"][0]["confirmed"] is False
    assert captured["response_format"]["type"] == "json_schema"
    assert captured["response_format"]["schema"]["additionalProperties"] is False


def test_ai_cannot_put_contacts_in_other_information(monkeypatch):
    from config.settings import LLMConfig

    text = "Nome profissional: Pessoa Exemplo\nTelefone: +55 00 00000-0000"
    monkeypatch.setattr(
        "core.ai_engine.invoke_llm",
        lambda *a, **k: json.dumps(
            {"entries": [{"category": "other", "text": text, "source_excerpt": text}]}
        ),
    )
    with pytest.raises(ValueError, match="misturou contatos"):
        intelligent_draft(text, "fixture.txt", LLMConfig())


@pytest.mark.parametrize("protocol", ["responses", "chat_completions"])
def test_json_schema_uses_protocol_specific_envelope(monkeypatch, protocol):
    from unittest.mock import Mock

    from config.settings import LLMConfig
    from core.ai_engine import invoke_llm

    response = Mock(status_code=200, headers={})
    response.json.return_value = (
        {"output": [{"type": "message", "content": [{"type": "output_text", "text": "{}"}]}]}
        if protocol == "responses"
        else {"choices": [{"message": {"content": "{}"}}]}
    )
    post = Mock(return_value=response)
    monkeypatch.setattr("core.ai_engine.requests.post", post)
    schema = {
        "type": "json_schema",
        "name": "profile",
        "schema": {"type": "object", "properties": {}, "additionalProperties": False},
        "strict": True,
    }
    invoke_llm(
        "Return JSON",
        LLMConfig(
            provider="openai", model="gpt-6.1-sol", api_key="fixture-only", protocol=protocol
        ),
        response_format=schema,
    )
    body = post.call_args.kwargs["json"]
    if protocol == "responses":
        assert body["text"]["format"] == schema
    else:
        assert body["response_format"]["json_schema"]["schema"] == schema["schema"]
        assert "schema" not in body["response_format"]


def test_selection_does_not_send_private_source_notes_or_contacts(document, monkeypatch):
    facts = draft_profile(document, "fixture.md")["entries"]
    captured = []
    monkeypatch.setattr(
        "core.ai_engine.invoke_llm",
        lambda prompt, *a, **k: captured.append(prompt) or json.dumps([facts[0]["id"]]),
    )
    assert use_ai_selection(facts, {"description": "Python"}, None)
    assert "Projeto Sigiloso" not in captured[0] and "source_excerpt" not in captured[0]
    assert "pessoa@example.invalid" not in captured[0]
